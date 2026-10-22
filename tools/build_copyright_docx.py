# -*- coding: utf-8 -*-
"""生成软著提交用 Word 文档。

输出：
    copyright/源程序.docx        每页 50 行，页眉为软件名+版本，页脚为页码
    copyright/软件说明书.docx    封面 + 由 docs 下 Markdown 合并而成

用法：
    python tools/build_copyright_docx.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
NAME = "面料显微图像双模型智能识别软件"
VERSION = "V1.0.0"
LINES_PER_PAGE = 50
MANUAL_PARTS = ("USER_MANUAL.md", "SOFTWARE_DESIGN.md", "ALGORITHM.md", "TEST_REPORT.md")


def _rfonts(parent_style, latin: str, east_asia: str) -> None:
    rpr = parent_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def setup_doc(header_text: str, font_size: int = 9) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(header_text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for piece in ("第 ", "PAGE", " 页 / 共 ", "NUMPAGES", " 页"):
        if piece in ("PAGE", "NUMPAGES"):
            _add_field(footer, piece)
        else:
            footer.add_run(piece).font.size = Pt(9)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(font_size)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    _rfonts(normal, "Consolas", "宋体")
    return doc


def build_source_doc() -> Path:
    lines = (ROOT / "copyright" / "source_full.txt").read_text(encoding="utf-8").splitlines()
    lines = [ln for ln in lines if not ln.startswith("// ---------- 第")]
    doc = setup_doc(f"{NAME} {VERSION}", font_size=9)
    for i, line in enumerate(lines):
        payload = line.rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(payload if payload else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if (i + 1) % LINES_PER_PAGE == 0 and i != len(lines) - 1:
            p.add_run().add_break(WD_BREAK.PAGE)
    out = ROOT / "copyright" / "源程序.docx"
    doc.save(out)
    return out


def _flush_table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(set(x) <= set("-: ") for x in c)]
    if not cells:
        return
    ncol = max(len(c) for c in cells)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Table Grid"
    for row in cells:
        target = table.add_row().cells
        for j in range(ncol):
            target[j].text = row[j] if j < len(row) else ""


def _add_markdown(doc: Document, path: Path) -> None:
    in_code = False
    table: list[str] = []
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            continue
        if line.startswith("|"):
            table.append(line)
            continue
        if table:
            _flush_table(doc, table)
            table = []
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("!["):
            doc.add_paragraph(f"【此处插入截图：{line.strip('![]()')}】")
        else:
            doc.add_paragraph(line)
    if table:
        _flush_table(doc, table)


def build_manual_doc() -> Path:
    doc = setup_doc(f"{NAME} {VERSION}", font_size=10.5)
    for _ in range(3):
        doc.add_paragraph()
    for text, size, bold in ((NAME, 24, True), ("软件说明书", 20, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
    for _ in range(6):
        doc.add_paragraph()
    for text in (f"版本：{VERSION}", "（软件著作权登记用）"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(text).font.size = Pt(14)
    doc.add_page_break()

    for name in MANUAL_PARTS:
        path = ROOT / "docs" / name
        if path.is_file():
            _add_markdown(doc, path)
            doc.add_page_break()

    out = ROOT / "copyright" / "软件说明书.docx"
    doc.save(out)
    return out


def main() -> None:
    src = build_source_doc()
    manual = build_manual_doc()
    print(f"已生成: {src.relative_to(ROOT)}")
    print(f"已生成: {manual.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
